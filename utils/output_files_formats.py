"""
Google Maps Output Writers
==========================

Purpose:
    Append extracted Maps records to the user-selected storage format while
    serializing writes from concurrent scraper workers.

Pipeline:
    google_maps_scraper.py -> output_files_formats.py -> google_maps_data.*

Input:
    A list of business dictionaries and a shared thread lock.

Output:
    ``google_maps_data.csv``, ``google_maps_data.json``, or
    ``google_maps_data.xlsx``. The CSV variant normally feeds ``build_leads.py``.
"""

from openpyxl import Workbook, load_workbook
from threading import Lock
from csv import DictReader, DictWriter
from tempfile import NamedTemporaryFile
import json
import os


class CSVCreator:
    """Append Maps records to a UTF-8 CSV, creating its header when needed."""

    def __init__(self, file_lock: Lock, output_path: str = "./CSV_FILES"):
        self._output_path = output_path
        self._file_lock = file_lock

    def create(self, list_of_dict_data: list[dict]):
        with self._file_lock:
            os.makedirs(self._output_path, exist_ok=True)
            file_name = "google_maps_data.csv"
            _isheader_file = False
            if not os.path.isfile(self._output_path + "/" + file_name):
                _isheader_file = True

            fieldnames = list(list_of_dict_data[0].keys())
            if not _isheader_file:
                file_path = self._output_path + "/" + file_name
                with open(file_path, "r", newline="", encoding="utf-8-sig") as existing_file:
                    reader = DictReader(existing_file)
                    existing_fields = list(reader.fieldnames or ())
                    existing_rows = list(reader)
                added_fields = [field for field in fieldnames if field not in existing_fields]
                if added_fields:
                    # Older Maps CSVs need a one-time schema migration before
                    # appending, otherwise DictWriter would silently drop the
                    # newly collected geography fields.
                    migrated_fields = [*existing_fields, *added_fields]
                    temporary_name = None
                    try:
                        with NamedTemporaryFile(
                            "w", newline="", encoding="utf-8-sig",
                            dir=self._output_path, prefix=f".{file_name}.",
                            suffix=".tmp", delete=False,
                        ) as temporary_file:
                            temporary_name = temporary_file.name
                            migrated_writer = DictWriter(
                                temporary_file, fieldnames=migrated_fields,
                                extrasaction="ignore",
                            )
                            migrated_writer.writeheader()
                            migrated_writer.writerows(existing_rows)
                        os.replace(temporary_name, file_path)
                        temporary_name = None
                    finally:
                        if temporary_name and os.path.exists(temporary_name):
                            os.unlink(temporary_name)
                    fieldnames = migrated_fields
                else:
                    fieldnames = existing_fields

            if _isheader_file:
                file_handler = open(self._output_path + "/" + file_name, "w", newline="", encoding="utf-8-sig")
            else:
                file_handler = open(self._output_path + "/" + file_name, "a", newline="", encoding="utf-8-sig")

            writer = DictWriter(file_handler, fieldnames=fieldnames, extrasaction='ignore')
            if _isheader_file:
                writer.writeheader()

            writer.writerows(list_of_dict_data)
            file_handler.close()


class JSONCreator:
    """Append Maps records to a JSON array under a shared writer lock."""

    def __init__(self, file_lock: Lock, output_path: str = "./JSON_FILES"):
        self._file_lock = file_lock
        self._output_path = output_path

    def create(self, list_of_dict_data: list[dict]):
        file_name = "google_maps_data.json"
        file_path = os.path.join(self._output_path, file_name)

        with self._file_lock:
            os.makedirs(self._output_path, exist_ok=True)
            if not os.path.isfile(file_path):
                with open(file_path, 'w', encoding='utf-8') as f:
                    json.dump(list_of_dict_data, f, ensure_ascii=False, indent=4)
            else:
                with open(file_path, 'r+', encoding='utf-8') as f:
                    try:
                        existing_data = json.load(f)
                    except json.JSONDecodeError:
                        existing_data = []
                    existing_data.extend(list_of_dict_data)
                    f.seek(0)
                    f.truncate()
                    json.dump(existing_data, f, ensure_ascii=False, indent=4)



class XLSXCreator:
    """Append Maps records to an Excel worksheet with stable column order."""

    def __init__(self, file_lock: Lock, output_path: str = "./XLSX_FILES"):
        self._file_lock = file_lock
        self._output_path = output_path

    def create(self, list_of_dict_data: list[dict]):
        file_name = "google_maps_data.xlsx"
        file_path = os.path.join(self._output_path, file_name)

        with self._file_lock:
            os.makedirs(self._output_path, exist_ok=True)
            if not os.path.isfile(file_path):
                wb = Workbook()
                ws = wb.active
                ws.title = "Data"
                headers = list(list_of_dict_data[0].keys())
                ws.append(headers)
                for data_dict in list_of_dict_data:
                    row_values = [data_dict.get(h, "") for h in headers]
                    ws.append(row_values)

                wb.save(file_path)
            else:
                wb = load_workbook(file_path)
                ws = wb.active
                headers = list(ws.iter_rows(min_row=1, max_row=1, values_only=True))[0]
                for data_dict in list_of_dict_data:
                    row_values = [data_dict.get(h, "") for h in headers]
                    ws.append(row_values)
                wb.save(file_path)


if __name__ == '__main__':
    data = [
        {"Name": "Store A", "Address": "123 Main St", "Rating": 4.5},
        {"Name": "Store B", "Address": "456 Elm St", "Rating": 3.7}
    ]

    file_thread_lock = Lock()
    csv_creator = CSVCreator(file_thread_lock, output_path="./CSV_FILES")
    csv_creator.create(data)

    # JSON usage
    json_creator = JSONCreator(file_thread_lock, output_path="./JSON_FILES")
    json_creator.create(data)

    # XLSX usage
    xlsx_creator = XLSXCreator(file_thread_lock, output_path="./XLSX_FILES")
    xlsx_creator.create(data)


