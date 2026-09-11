"""Coverage for persistent discovery timestamps and daily exports."""

from csv import DictReader, DictWriter
from datetime import date
from pathlib import Path
from tempfile import TemporaryDirectory
from threading import Lock
from unittest import TestCase
from unittest.mock import patch

from openpyxl import load_workbook

from utils.build_leads import MASTER_FIELDS, build_lead_files
from utils.build_outreach import OUTPUT_FIELDS, build_outreach
from utils.build_outreach_queue import build_queue, read_csv as read_queue_csv
from utils.enrich_leads import enrich_leads_batch
from utils.export_daily_leads import export_daily_leads
from utils.enrich_search_emails import prepare_records as prepare_email_records
from utils.google_maps_scraper import GoogleMaps
from utils.google_search_discovery import (
    DISCOVERY_FIELDS,
    discover_companies,
    merge_accepted_discoveries,
    merge_discoveries,
)
from utils.known_companies import KnownCompanies
from utils.output_files_formats import CSVCreator


STAMP_1 = "2026-09-01T10:30:00+01:00"
STAMP_2 = "2026-09-03T14:00:00+01:00"


def write_csv(path, fields, rows):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def read_csv(path):
    with path.open("r", newline="", encoding="utf-8-sig") as handle:
        return list(DictReader(handle))


def maps_row(name="Acme", stamp=STAMP_1):
    return {
        "title": name, "webpage": "https://acme.test", "phone_number": "+216 1",
        "site_email": "info@acme.test", "added_at": stamp,
    }


def search_row(stamp=STAMP_2):
    return {
        "company_name": "Acme Search", "website": "https://acme.test/",
        "source": "google_search", "source_query": "software Tunisia",
        "source_url": "https://acme.test/about", "country": "", "city": "",
        "location": "", "added_at": stamp,
    }


class FakeDriver:
    current_window_handle = "main"

    def quit(self):
        pass


class DiscoveryTimestampTests(TestCase):
    def test_new_maps_company_receives_timestamp_and_duplicate_does_not_append(self):
        registry = KnownCompanies()
        scraper = GoogleMaps(incremental=True, known_companies=registry, verbose=False)
        scraper.validate_result_link = lambda *args: ("1", "2", "https://maps/X")
        scraper.get_title = lambda driver: "New Co"
        scraper.get_website_link = lambda driver: "https://new.test"
        scraper.get_phone_number = lambda driver: "+216 2"
        scraper.get_cover_image = lambda driver: ""
        scraper.get_rating_in_card = lambda driver: ""
        scraper.get_privacy_price = lambda driver: ""
        scraper.get_category = lambda driver: ""
        scraper.get_address = lambda driver: ""
        scraper.get_working_hours = lambda driver: ""
        scraper.get_menu_link = lambda driver: ""
        scraper.get_related_images_list = lambda driver: ""
        scraper.get_about_description = lambda driver: {}
        scraper.reset_driver_for_next_run = lambda *args: None
        scraper._web_pattern_scraper.find_patterns = lambda *args: {}
        stored = []
        scraper._file_creator.create = lambda list_of_dict_data: stored.extend(list_of_dict_data)
        with patch("utils.google_maps_scraper.discovery_timestamp", return_value=STAMP_1):
            self.assertEqual(scraper._scrape_result_and_store(FakeDriver(), "continue", "q", [1, 1]), "new")
            self.assertEqual(scraper._scrape_result_and_store(FakeDriver(), "continue", "q", [1, 1]), "same_run")
        self.assertEqual([row["added_at"] for row in stored], [STAMP_1])

    def test_new_search_company_receives_timestamp(self):
        with patch("utils.google_search_discovery.discovery_timestamp", return_value=STAMP_1):
            rows = merge_accepted_discoveries([], [search_row(stamp="")])
        self.assertEqual(rows[0]["added_at"], STAMP_1)

    def test_known_legacy_and_same_run_search_duplicates_get_no_new_timestamp(self):
        legacy = search_row(stamp="")
        with patch("utils.google_search_discovery.discovery_timestamp", return_value=STAMP_1) as clock:
            known = merge_accepted_discoveries([legacy], [search_row(stamp="")])
            same_run = merge_accepted_discoveries([], [search_row(stamp=""), search_row(stamp="")])
        self.assertEqual(known[0]["added_at"], "")
        self.assertEqual(same_run[0]["added_at"], STAMP_1)
        self.assertEqual(clock.call_count, 1)

    def test_search_incremental_known_duplicate_is_not_timestamped(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            query = root / "queries.txt"
            query.write_text("software\n", encoding="utf-8")
            output = root / "google_search_companies.csv"
            write_csv(output, DISCOVERY_FIELDS, [search_row(stamp="")])
            with patch("utils.google_search_discovery.search_query", return_value=([search_row(stamp="")], "")), patch(
                "utils.google_search_discovery.discovery_timestamp"
            ) as clock:
                discover_companies(query, output, limit=1, delay=0, incremental=True,
                                   driver_factory=lambda windowed=False: FakeDriver())
            self.assertEqual(read_csv(output)[0]["added_at"], "")
            clock.assert_not_called()

    def test_old_maps_csv_migrates_without_changing_rows(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            path = root / "google_maps_data.csv"
            write_csv(path, ("title", "webpage"), [{"title": "Old", "webpage": "https://old.test"}])
            CSVCreator(Lock(), str(root)).create([maps_row("New")])
            rows = read_csv(path)
            self.assertEqual(rows[0]["title"], "Old")
            self.assertEqual(rows[0]["added_at"], "")
            self.assertEqual(rows[1]["added_at"], STAMP_1)


class BuildTimestampTests(TestCase):
    def build(self, maps_rows, search_rows=None, existing=None):
        temporary = TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        root = Path(temporary.name)
        write_csv(root / "google_maps_data.csv", tuple(maps_rows[0]) if maps_rows else ("title", "webpage", "added_at"), maps_rows)
        if search_rows is not None:
            write_csv(root / "google_search_companies.csv", DISCOVERY_FIELDS, search_rows)
        if existing is not None:
            write_csv(root / "leads_master.csv", MASTER_FIELDS, existing)
        build_lead_files(root / "google_maps_data.csv", root)
        return root, read_csv(root / "leads_master.csv")

    @staticmethod
    def master_row(stamp=""):
        row = {field: "" for field in MASTER_FIELDS}
        row.update({
            "name": "Acme", "website": "https://acme.test",
            "added_at": stamp,
        })
        return row

    def test_old_master_blank_and_maps_valid_keeps_maps_timestamp(self):
        _, rows = self.build([maps_row(stamp=STAMP_1)], existing=[self.master_row()])
        self.assertEqual(rows[0]["added_at"], STAMP_1)

    def test_old_master_valid_and_maps_blank_keeps_master_timestamp(self):
        _, rows = self.build([maps_row(stamp="")], existing=[self.master_row(STAMP_1)])
        self.assertEqual(rows[0]["added_at"], STAMP_1)

    def test_old_master_later_and_maps_earlier_keeps_earliest_timestamp(self):
        _, rows = self.build(
            [maps_row(stamp=STAMP_1)], existing=[self.master_row(STAMP_2)],
        )
        self.assertEqual(rows[0]["added_at"], STAMP_1)

    def test_maps_later_and_search_earlier_keeps_earliest_timestamp(self):
        _, rows = self.build([maps_row(stamp=STAMP_2)], [search_row(STAMP_1)])
        self.assertEqual(rows[0]["added_at"], STAMP_1)

    def test_all_timestamp_sources_blank_remain_blank(self):
        _, rows = self.build(
            [maps_row(stamp="")], [search_row("")], [self.master_row()],
        )
        self.assertEqual(rows[0]["added_at"], "")

    def test_rebuild_is_idempotent(self):
        root, _ = self.build(
            [maps_row(stamp=STAMP_2)], [search_row(STAMP_1)],
            [self.master_row()],
        )
        output_names = ("leads_master.csv", "leads_ready.csv", "leads_review.csv")
        first_build = {name: (root / name).read_bytes() for name in output_names}

        build_lead_files(root / "google_maps_data.csv", root)

        self.assertEqual(
            {name: (root / name).read_bytes() for name in output_names},
            first_build,
        )

    def test_legacy_row_without_timestamp_evidence_remains_blank(self):
        _, rows = self.build([maps_row(stamp="")])
        self.assertEqual(rows[0]["added_at"], "")

    def test_source_timestamps_propagate_to_ready_and_review_outputs(self):
        ready = maps_row(name="Ready", stamp=STAMP_1)
        review = {
            **maps_row(name="Review", stamp=STAMP_2),
            "webpage": "https://review.test",
            "site_email": "contact@different.test",
        }
        root, master_rows = self.build([ready, review])

        self.assertEqual(
            {row["name"]: row["added_at"] for row in master_rows},
            {"Ready": STAMP_1, "Review": STAMP_2},
        )
        self.assertEqual(read_csv(root / "leads_ready.csv")[0]["added_at"], STAMP_1)
        self.assertEqual(read_csv(root / "leads_review.csv")[0]["added_at"], STAMP_2)


class PropagationTests(TestCase):
    def test_email_enrichment_preserves_timestamp(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            master = {field: "" for field in MASTER_FIELDS}
            master.update({
                "name": "Search Co", "website": "https://search.test",
                "source": "google_search", "added_at": STAMP_1,
            })
            write_csv(root / "master.csv", MASTER_FIELDS, [master])
            records = prepare_email_records(root / "master.csv", root / "emails.csv")
            self.assertEqual(records[0]["added_at"], STAMP_1)

    def test_company_enrichment_preserves_timestamp(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            ready = root / "leads_ready.csv"
            output = root / "leads_enriched.csv"
            write_csv(ready, ("name", "email", "phone", "website", "added_at"), [{
                "name": "Acme", "email": "info@acme.test", "phone": "", "website": "https://acme.test", "added_at": STAMP_1,
            }])
            enrich_leads_batch(ready, output, limit=1, driver_factory=FakeDriver,
                               enrichment_function=lambda *args, **kwargs: {"description": "Acme software", "industry": "Software", "services": "Development"})
            self.assertEqual(read_csv(output)[0]["added_at"], STAMP_1)

    def test_outreach_and_queue_preserve_timestamp(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            enriched = {field: "" for field in OUTPUT_FIELDS}
            enriched.update({"company_name": "Acme", "email": "info@acme.test", "website": "https://acme.test", "description": "Acme software", "industry": "Software", "services": "Development", "enrichment_status": "SUCCESS", "added_at": STAMP_1})
            write_csv(root / "leads_enriched_final.csv", OUTPUT_FIELDS, [enriched])
            build_outreach(root / "leads_enriched_final.csv", root / "outreach_ready.csv", root / "outreach_review.csv")
            build_queue(root / "outreach_ready.csv", root / "history.csv", root / "outreach_queue.csv")
            self.assertEqual(read_csv(root / "outreach_ready.csv")[0]["added_at"], STAMP_1)
            self.assertEqual(read_queue_csv(root / "outreach_queue.csv")[1][0]["added_at"], STAMP_1)


class DailyExportTests(TestCase):
    def test_filters_blanks_invalid_other_dates_and_converts_timezone(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            fields = ("name", "website", "added_at")
            write_csv(root / "master.csv", fields, [
                {"name": "Match", "website": "https://a.test", "added_at": STAMP_1},
                {"name": "UTC midnight", "website": "https://b.test", "added_at": "2026-08-31T23:30:00+00:00"},
                {"name": "Other", "website": "https://c.test", "added_at": STAMP_2},
                {"name": "Legacy", "website": "https://d.test", "added_at": ""},
                {"name": "Broken", "website": "https://e.test", "added_at": "not-a-date"},
            ])
            summary = export_daily_leads(root / "master.csv", date(2026, 9, 1), root / "exports")
            csv_rows = read_csv(summary["csv"])
            workbook = load_workbook(summary["xlsx"], read_only=True)
            xlsx_names = [row[0] for row in list(workbook.active.iter_rows(values_only=True))[1:]]
            self.assertEqual([row["name"] for row in csv_rows], ["Match", "UTC midnight"])
            self.assertEqual(xlsx_names, ["Match", "UTC midnight"])
            self.assertEqual((summary["blank"], summary["invalid"]), (1, 1))


class MergeTimestampTests(TestCase):
    def test_invalid_timestamp_does_not_replace_valid_earlier_timestamp(self):
        rows = merge_discoveries([], [search_row(STAMP_2), search_row("invalid")])
        self.assertEqual(rows[0]["added_at"], STAMP_2)
