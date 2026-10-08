"""Conservative, auditable backfill of historical company data.

Historical timestamps are observation evidence only.  This importer never sets
``first_seen_at`` and never classifies a historical entity as ``NEW``.
"""

from __future__ import annotations

from argparse import ArgumentParser
from collections import defaultdict
from dataclasses import asdict, dataclass, field
from datetime import datetime, timezone
import csv
from hashlib import sha256
import json
from pathlib import Path
import sqlite3
from tempfile import TemporaryDirectory
from typing import Iterable, Iterator
from urllib.parse import urlsplit
from uuid import UUID, uuid5

from company_registry.storage import DEFAULT_DATABASE, connect_registry
from utils.build_leads import (
    FREE_EMAIL_DOMAINS,
    clean_value,
    normalize_name,
    normalize_phone,
)
from utils.maps_identity import maps_identities_for


REGISTRY_NAMESPACE = UUID("c5c34fab-4f32-53b7-a268-2cd95b468fe3")
SUPPORTED_SUFFIXES = {".csv", ".xlsx"}


@dataclass(frozen=True)
class SourceFile:
    path: Path
    relative_path: str
    kind: str


@dataclass
class ImportReport:
    dry_run: bool
    database: str
    run_id: str = ""
    sources_inspected: list[str] = field(default_factory=list)
    sources_loaded: list[str] = field(default_factory=list)
    sources_skipped_as_duplicates: list[str] = field(default_factory=list)
    records_examined: int = 0
    records_imported: int = 0
    exact_place_id_matches: int = 0
    supporting_identity_matches: int = 0
    ambiguous_records: int = 0
    quarantined_records: int = 0
    proposed_companies: int = 0
    proposed_branches: int = 0
    existing_companies: int = 0
    existing_branches: int = 0
    reliable_observation_timestamps: int = 0
    backup_path: str | None = None
    resolution_reasons: dict[str, int] = field(default_factory=dict)

    def to_json(self) -> str:
        return json.dumps(asdict(self), indent=2, sort_keys=True)


@dataclass(frozen=True)
class HistoricalRecord:
    source: SourceFile
    row_number: int
    raw: dict[str, str]
    content_hash: str
    name: str
    name_key: str
    place_id: str
    website: str
    website_domain: str
    email_domain: str
    phone: str
    phone_key: str
    address: str
    address_key: str
    observed_at: str | None

    @property
    def company_domain(self) -> str:
        return self.website_domain or self.email_domain


@dataclass
class BranchState:
    branch_id: str
    company_id: str
    place_id: str = ""
    names: set[str] = field(default_factory=set)
    domains: set[str] = field(default_factory=set)
    phones: set[str] = field(default_factory=set)
    addresses: set[str] = field(default_factory=set)


def _stable_uuid(kind: str, value: str) -> str:
    return str(uuid5(REGISTRY_NAMESPACE, f"{kind}|{value}"))


def _canonical_json(row: dict[str, object]) -> str:
    normalized = {
        str(key): "" if value is None else str(value).strip()
        for key, value in row.items()
    }
    return json.dumps(normalized, ensure_ascii=False, sort_keys=True, separators=(",", ":"))


def _domain(value: str) -> str:
    value = clean_value(value)
    if not value:
        return ""
    parsed = urlsplit(value if "://" in value else "//" + value)
    domain = (parsed.hostname or "").casefold().rstrip(".")
    return domain[4:] if domain.startswith("www.") else domain


def _email_domain(row: dict[str, str]) -> str:
    values = " ".join((
        row.get("site_email", ""), row.get("email", ""),
        row.get("alternate_emails", ""),
    ))
    domains = {
        value.rsplit("@", 1)[1].casefold().strip(" .,;<>[]()")
        for value in values.replace(";", " ").replace(",", " ").split()
        if "@" in value
    }
    domains = {value for value in domains if "." in value and value not in FREE_EMAIL_DOMAINS}
    return sorted(domains)[0] if len(domains) == 1 else ""


def _timestamp(value: str) -> str | None:
    value = clean_value(value)
    if not value:
        return None
    try:
        parsed = datetime.fromisoformat(value.replace("Z", "+00:00"))
    except ValueError:
        return None
    if parsed.tzinfo is None:
        return None
    return parsed.isoformat()


def _source_kind(relative_path: str) -> str:
    name = Path(relative_path).name.casefold()
    if "google_maps" in name:
        return "GOOGLE_MAPS"
    if "google_search" in name or "search_email" in name:
        return "GOOGLE_SEARCH"
    if "daily_exports" in relative_path:
        return "DAILY_EXPORT"
    if "enrich" in name:
        return "ENRICHMENT"
    if "outreach" in name:
        return "OUTREACH"
    if "lead" in name:
        return "LEAD_EXPORT"
    return "OTHER"


def discover_sources(project_root: Path) -> tuple[list[SourceFile], list[str], list[str]]:
    """Return canonical input files plus inspected and duplicate-file lists."""
    data_root = project_root / "CSV_FILES"
    candidates = sorted(
        path for path in data_root.rglob("*")
        if path.is_file() and path.suffix.casefold() in SUPPORTED_SUFFIXES
        and not path.name.casefold().startswith("google_jobs")
    )
    inspected = [path.relative_to(project_root).as_posix() for path in candidates]
    csv_stems = {path.with_suffix("") for path in candidates if path.suffix.casefold() == ".csv"}
    seen_hashes: dict[str, Path] = {}
    selected: list[SourceFile] = []
    skipped: list[str] = []
    for path in candidates:
        relative = path.relative_to(project_root).as_posix()
        if path.suffix.casefold() == ".xlsx" and path.with_suffix("") in csv_stems:
            skipped.append(relative + " (CSV mirror)")
            continue
        digest = sha256(path.read_bytes()).hexdigest()
        if digest in seen_hashes:
            original = seen_hashes[digest].relative_to(project_root).as_posix()
            skipped.append(relative + f" (byte-identical to {original})")
            continue
        seen_hashes[digest] = path
        selected.append(SourceFile(path, relative, _source_kind(relative)))
    return selected, inspected, skipped


def _read_rows(source: SourceFile) -> Iterator[tuple[int, dict[str, str]]]:
    if source.path.suffix.casefold() == ".csv":
        with source.path.open("r", newline="", encoding="utf-8-sig", errors="replace") as handle:
            for row_number, row in enumerate(csv.DictReader(handle), 2):
                yield row_number, {key: value or "" for key, value in row.items() if key is not None}
        return
    try:
        from openpyxl import load_workbook
    except ImportError as error:  # pragma: no cover - dependency is declared by the project
        raise RuntimeError("XLSX history requires the declared openpyxl dependency") from error
    workbook = load_workbook(source.path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    headers = [str(value or "") for value in next(rows, ())]
    for row_number, values in enumerate(rows, 2):
        yield row_number, {
            key: "" if value is None else str(value)
            for key, value in zip(headers, values)
        }
    workbook.close()


def load_records(sources: Iterable[SourceFile]) -> Iterator[HistoricalRecord]:
    for source in sources:
        for row_number, row in _read_rows(source):
            raw_json = _canonical_json(row)
            name = clean_value(row.get("title") or row.get("company_name") or row.get("name"))
            website = clean_value(row.get("webpage") or row.get("website") or row.get("source_url"))
            phone = clean_value(row.get("phone_number") or row.get("phone"))
            address = clean_value(row.get("address") or row.get("location"))
            yield HistoricalRecord(
                source=source,
                row_number=row_number,
                raw=row,
                content_hash=sha256(raw_json.encode("utf-8")).hexdigest(),
                name=name,
                name_key=normalize_name(name),
                place_id=maps_identities_for(row)["place_id"],
                website=website,
                website_domain=_domain(website),
                email_domain=_email_domain(row),
                phone=phone,
                phone_key=normalize_phone(phone),
                address=address,
                address_key=normalize_name(address),
                observed_at=_timestamp(row.get("added_at", "")),
            )


class Resolver:
    def __init__(self, connection: sqlite3.Connection, now: str):
        self.connection = connection
        self.now = now
        self.branches: dict[str, BranchState] = {}
        self.place_index: dict[str, str] = {}
        self.evidence_indexes: dict[str, dict[tuple[str, str], set[str]]] = {
            "phone": defaultdict(set), "address": defaultdict(set), "domain": defaultdict(set),
        }
        self.company_index: dict[tuple[str, str], set[str]] = defaultdict(set)
        self.source_resolutions: dict[tuple[str, int, str], sqlite3.Row] = {}
        self._load_existing()

    def _load_existing(self) -> None:
        identities: dict[str, dict[str, set[str]]] = defaultdict(lambda: defaultdict(set))
        for row in self.connection.execute(
            "SELECT company_id, identity_type, normalized_value FROM company_identities"
        ):
            identities[row["company_id"]][row["identity_type"]].add(row["normalized_value"])
        for company_id, values in identities.items():
            for name in values["NAME"]:
                for domain in values["WEBSITE_DOMAIN"] | values["EMAIL_DOMAIN"]:
                    self.company_index[(name, domain)].add(company_id)
        for row in self.connection.execute("SELECT * FROM branches"):
            state = BranchState(
                branch_id=row["branch_id"], company_id=row["company_id"],
                place_id=row["google_maps_place_id"] or "",
                names={normalize_name(row["display_name"] or "")} - {""},
                domains={_domain(row["website_url"] or "")} - {""},
                phones={normalize_phone(row["phone"] or "")} - {""},
                addresses={normalize_name(row["address"] or "")} - {""},
            )
            self.branches[state.branch_id] = state

        for row in self.connection.execute(
            """SELECT source_path, source_row_number, source_content_hash,
                      resolution_status, company_id, branch_id, reason,
                      raw_record_json
                 FROM historical_source_records"""
        ):
            key = (row["source_path"], row["source_row_number"], row["source_content_hash"])
            self.source_resolutions[key] = row
            if row["resolution_status"] != "IMPORTED" or not row["branch_id"]:
                continue
            state = self.branches.get(row["branch_id"])
            if state is None:
                continue
            raw = json.loads(row["raw_record_json"])
            name = clean_value(raw.get("title") or raw.get("company_name") or raw.get("name"))
            website = clean_value(raw.get("webpage") or raw.get("website") or raw.get("source_url"))
            phone = clean_value(raw.get("phone_number") or raw.get("phone"))
            address = clean_value(raw.get("address") or raw.get("location"))
            domain = _domain(website) or _email_domain(raw)
            state.names.update({normalize_name(name)} - {""})
            state.domains.update({domain} - {""})
            state.phones.update({normalize_phone(phone)} - {""})
            state.addresses.update({normalize_name(address)} - {""})

        for state in self.branches.values():
            self._index(state)

    def prior_resolution(self, record: HistoricalRecord) -> sqlite3.Row | None:
        """Return the immutable audit decision for an already-processed source row."""
        return self.source_resolutions.get((
            record.source.relative_path, record.row_number, record.content_hash,
        ))

    def _index(self, state: BranchState) -> None:
        self.branches[state.branch_id] = state
        if state.place_id:
            self.place_index[state.place_id] = state.branch_id
        for name in state.names:
            for value in state.phones:
                self.evidence_indexes["phone"][(name, value)].add(state.branch_id)
            for value in state.addresses:
                self.evidence_indexes["address"][(name, value)].add(state.branch_id)
            for value in state.domains:
                self.evidence_indexes["domain"][(name, value)].add(state.branch_id)

    def candidates(self, record: HistoricalRecord) -> set[str]:
        matches: set[str] = set()
        for kind, value in (
            ("phone", record.phone_key), ("address", record.address_key),
            ("domain", record.company_domain),
        ):
            if value:
                matches.update(self.evidence_indexes[kind].get((record.name_key, value), ()))
        return matches

    def company_for(self, record: HistoricalRecord, branch_key: str) -> tuple[str | None, str | None]:
        domain = record.company_domain
        candidates = self.company_index.get((record.name_key, domain), set()) if domain else set()
        if len(candidates) > 1:
            return None, "multiple companies share exact name and domain"
        if candidates:
            return next(iter(candidates)), None
        company_seed = f"{record.name_key}|{domain}" if domain else branch_key
        return _stable_uuid("company", company_seed), None

    def create(self, record: HistoricalRecord, branch_key: str) -> tuple[str | None, str | None, str | None]:
        company_id, error = self.company_for(record, branch_key)
        if error:
            return None, None, error
        branch_id = _stable_uuid("branch", branch_key)
        self.connection.execute(
            """INSERT OR IGNORE INTO companies
               (company_id, canonical_name, discovery_status, first_seen_at,
                last_seen_at, created_at, updated_at)
               VALUES (?, ?, 'LEGACY_UNKNOWN', NULL, ?, ?, ?)""",
            (company_id, record.name, record.observed_at, self.now, self.now),
        )
        self.connection.execute(
            """INSERT OR IGNORE INTO branches
               (branch_id, company_id, display_name, google_maps_place_id,
                website_url, phone, address, discovery_status, first_seen_at,
                last_seen_at, created_at, updated_at)
               VALUES (?, ?, ?, ?, ?, ?, ?, 'LEGACY_UNKNOWN', NULL, ?, ?, ?)""",
            (branch_id, company_id, record.name, record.place_id or None,
             record.website or None, record.phone or None, record.address or None,
             record.observed_at, self.now, self.now),
        )
        state = BranchState(
            branch_id, company_id, record.place_id, {record.name_key},
            {record.company_domain} - {""}, {record.phone_key} - {""},
            {record.address_key} - {""},
        )
        self._index(state)
        if record.company_domain:
            self.company_index[(record.name_key, record.company_domain)].add(company_id)
        return company_id, branch_id, None

    def enrich(self, record: HistoricalRecord, branch_id: str) -> str:
        state = self.branches[branch_id]
        state.names.add(record.name_key)
        state.domains.update({record.company_domain} - {""})
        state.phones.update({record.phone_key} - {""})
        state.addresses.update({record.address_key} - {""})
        self._index(state)
        if record.company_domain:
            self.company_index[(record.name_key, record.company_domain)].add(state.company_id)
        self.connection.execute(
            """UPDATE branches SET
                 display_name=COALESCE(display_name, ?),
                 website_url=COALESCE(website_url, ?), phone=COALESCE(phone, ?),
                 address=COALESCE(address, ?),
                 last_seen_at=CASE
                   WHEN ? IS NULL THEN last_seen_at
                   WHEN last_seen_at IS NULL OR last_seen_at < ? THEN ? ELSE last_seen_at END,
                 updated_at=? WHERE branch_id=?""",
            (record.name or None, record.website or None, record.phone or None,
             record.address or None, record.observed_at, record.observed_at,
             record.observed_at, self.now, branch_id),
        )
        self.connection.execute(
            """UPDATE companies SET canonical_name=COALESCE(canonical_name, ?),
                 last_seen_at=CASE
                   WHEN ? IS NULL THEN last_seen_at
                   WHEN last_seen_at IS NULL OR last_seen_at < ? THEN ? ELSE last_seen_at END,
                 updated_at=? WHERE company_id=?""",
            (record.name or None, record.observed_at, record.observed_at,
             record.observed_at, self.now, state.company_id),
        )
        return state.company_id


def _identity(connection: sqlite3.Connection, company_id: str, kind: str,
              raw: str, normalized: str, observed_at: str | None, now: str) -> None:
    if not normalized:
        return
    identity_id = _stable_uuid("identity", f"{company_id}|{kind}|{normalized}")
    connection.execute(
        """INSERT INTO company_identities
           (identity_id, company_id, identity_type, identity_value,
            normalized_value, first_seen_at, last_seen_at, created_at, updated_at)
           VALUES (?, ?, ?, ?, ?, NULL, ?, ?, ?)
           ON CONFLICT(company_id, identity_type, normalized_value) DO UPDATE SET
             last_seen_at=CASE
               WHEN excluded.last_seen_at IS NULL THEN company_identities.last_seen_at
               WHEN company_identities.last_seen_at IS NULL
                 OR company_identities.last_seen_at < excluded.last_seen_at
               THEN excluded.last_seen_at ELSE company_identities.last_seen_at END,
             updated_at=excluded.updated_at""",
        (identity_id, company_id, kind, raw, normalized, observed_at, now, now),
    )


def _record_provenance(connection: sqlite3.Connection, record: HistoricalRecord,
                       run_id: str, status: str, now: str,
                       company_id: str | None = None, branch_id: str | None = None,
                       reason: str | None = None) -> None:
    source_record_id = _stable_uuid(
        "source-record",
        f"{record.source.relative_path}|{record.row_number}|{record.content_hash}",
    )
    connection.execute(
        """INSERT OR IGNORE INTO historical_source_records
           (source_record_id, run_id, source_path, source_kind,
            source_row_number, source_content_hash, resolution_status,
            company_id, branch_id, observed_at, reason, raw_record_json, created_at)
           VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
        (source_record_id, run_id, record.source.relative_path, record.source.kind,
         record.row_number, record.content_hash, status, company_id, branch_id,
         record.observed_at, reason, _canonical_json(record.raw), now),
    )


def _backup(connection: sqlite3.Connection, database: Path) -> Path:
    stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%S%fZ")
    destination = database.with_name(f"{database.stem}.backup-{stamp}{database.suffix}")
    backup = sqlite3.connect(destination)
    try:
        connection.backup(backup)
    finally:
        backup.close()
    return destination


def _import_into(connection: sqlite3.Connection, records: Iterable[HistoricalRecord],
                 report: ImportReport, *, fail_after_records: int | None = None) -> None:
    now = datetime.now(timezone.utc).isoformat()
    run_id = _stable_uuid("legacy-run", now)
    report.run_id = run_id
    report.existing_companies = connection.execute("SELECT count(*) FROM companies").fetchone()[0]
    report.existing_branches = connection.execute("SELECT count(*) FROM branches").fetchone()[0]
    initial_companies = report.existing_companies
    initial_branches = report.existing_branches
    reasons: dict[str, int] = defaultdict(int)
    connection.execute("BEGIN IMMEDIATE")
    try:
        connection.execute(
            """INSERT INTO discovery_runs
               (run_id, run_type, status, started_at, created_at)
               VALUES (?, 'LEGACY_IMPORT', 'RUNNING', ?, ?)""",
            (run_id, now, now),
        )
        resolver = Resolver(connection, now)
        observed_companies: set[str] = set()
        for record in records:
            report.records_examined += 1
            if record.observed_at:
                report.reliable_observation_timestamps += 1
            prior = resolver.prior_resolution(record)
            if prior is not None:
                status = prior["resolution_status"]
                reason = prior["reason"]
                if status == "IMPORTED":
                    report.records_imported += 1
                    observed_companies.add(prior["company_id"])
                elif status == "AMBIGUOUS":
                    report.ambiguous_records += 1
                    reasons[reason or "previously ambiguous"] += 1
                else:
                    report.quarantined_records += 1
                    reasons[reason or "previously quarantined"] += 1
                continue
            if not record.name_key:
                reason = "missing company name"
                report.quarantined_records += 1
                reasons[reason] += 1
                _record_provenance(connection, record, run_id, "QUARANTINED", now, reason=reason)
                continue
            if not any((record.place_id, record.company_domain, record.phone_key, record.address_key)):
                reason = "insufficient identity evidence"
                report.quarantined_records += 1
                reasons[reason] += 1
                _record_provenance(connection, record, run_id, "QUARANTINED", now, reason=reason)
                continue

            branch_id: str | None = None
            company_id: str | None = None
            if record.place_id and record.place_id in resolver.place_index:
                branch_id = resolver.place_index[record.place_id]
                company_id = resolver.enrich(record, branch_id)
                report.exact_place_id_matches += 1
            elif record.place_id:
                company_id, branch_id, error = resolver.create(record, f"place|{record.place_id}")
                if error:
                    report.ambiguous_records += 1
                    reasons[error] += 1
                    _record_provenance(connection, record, run_id, "AMBIGUOUS", now, reason=error)
                    continue
            else:
                candidates = resolver.candidates(record)
                if len(candidates) > 1:
                    reason = "supporting identities match multiple branches"
                    report.ambiguous_records += 1
                    reasons[reason] += 1
                    _record_provenance(connection, record, run_id, "AMBIGUOUS", now, reason=reason)
                    continue
                if candidates:
                    branch_id = next(iter(candidates))
                    company_id = resolver.enrich(record, branch_id)
                    report.supporting_identity_matches += 1
                else:
                    if record.phone_key:
                        branch_key = f"name-phone|{record.name_key}|{record.phone_key}"
                    elif record.address_key:
                        branch_key = f"name-address|{record.name_key}|{record.address_key}"
                    else:
                        branch_key = f"name-domain|{record.name_key}|{record.company_domain}"
                    company_id, branch_id, error = resolver.create(record, branch_key)
                    if error:
                        report.ambiguous_records += 1
                        reasons[error] += 1
                        _record_provenance(connection, record, run_id, "AMBIGUOUS", now, reason=error)
                        continue

            assert company_id and branch_id
            _identity(connection, company_id, "NAME", record.name, record.name_key, record.observed_at, now)
            _identity(connection, company_id, "WEBSITE_DOMAIN", record.website_domain,
                      record.website_domain, record.observed_at, now)
            _identity(connection, company_id, "EMAIL_DOMAIN", record.email_domain,
                      record.email_domain, record.observed_at, now)
            observed_companies.add(company_id)
            report.records_imported += 1
            _record_provenance(
                connection, record, run_id, "IMPORTED", now,
                company_id=company_id, branch_id=branch_id,
            )
            if fail_after_records is not None and report.records_examined >= fail_after_records:
                raise RuntimeError("injected historical import failure")

        for company_id in observed_companies:
            connection.execute(
                """INSERT INTO run_companies
                   (run_id, company_id, discovery_status, observed_at, created_at)
                   VALUES (?, ?, 'LEGACY_UNKNOWN', ?, ?)""",
                (run_id, company_id, now, now),
            )
        final_status = "PARTIAL" if report.ambiguous_records or report.quarantined_records else "SUCCESS"
        connection.execute(
            "UPDATE discovery_runs SET status=?, finished_at=? WHERE run_id=?",
            (final_status, datetime.now(timezone.utc).isoformat(), run_id),
        )
        report.proposed_companies = connection.execute("SELECT count(*) FROM companies").fetchone()[0] - initial_companies
        report.proposed_branches = connection.execute("SELECT count(*) FROM branches").fetchone()[0] - initial_branches
        report.resolution_reasons = dict(sorted(reasons.items()))
        connection.commit()
    except BaseException:
        connection.rollback()
        raise


def import_history(project_root: Path, database: Path = DEFAULT_DATABASE, *,
                   dry_run: bool = True, sources: Iterable[SourceFile] | None = None,
                   fail_after_records: int | None = None) -> ImportReport:
    """Resolve historical files and either simulate or transactionally apply them."""
    project_root = Path(project_root).resolve()
    database = Path(database)
    if not database.is_absolute():
        database = project_root / database
    selected, inspected, skipped = discover_sources(project_root) if sources is None else (list(sources), [], [])
    report = ImportReport(
        dry_run=dry_run, database=str(database), sources_inspected=inspected,
        sources_loaded=[source.relative_path for source in selected],
        sources_skipped_as_duplicates=skipped,
    )

    if dry_run:
        with TemporaryDirectory(prefix="company-registry-dry-run-") as directory:
            simulation = Path(directory) / "company_registry.db"
            if database.exists():
                source_connection = sqlite3.connect(f"file:{database}?mode=ro", uri=True)
                target_connection = sqlite3.connect(simulation)
                try:
                    source_connection.backup(target_connection)
                finally:
                    source_connection.close()
                    target_connection.close()
            connection = connect_registry(simulation)
            try:
                _import_into(connection, load_records(selected), report,
                             fail_after_records=fail_after_records)
            finally:
                connection.close()
        return report

    connection = connect_registry(database)
    try:
        populated = sum(
            connection.execute(f"SELECT count(*) FROM {table}").fetchone()[0]
            for table in ("companies", "branches", "company_identities", "historical_source_records")
        )
        if populated:
            report.backup_path = str(_backup(connection, database))
        _import_into(connection, load_records(selected), report,
                     fail_after_records=fail_after_records)
    finally:
        connection.close()
    return report


def main() -> None:
    parser = ArgumentParser(description=__doc__)
    parser.add_argument("--project-root", type=Path, default=Path.cwd())
    parser.add_argument("--database", type=Path, default=DEFAULT_DATABASE)
    parser.add_argument(
        "--apply", action="store_true",
        help="Write the import. Without this flag a read-only production dry-run is used.",
    )
    parser.add_argument("--report", type=Path, help="Optional JSON report path")
    arguments = parser.parse_args()
    report = import_history(
        arguments.project_root, arguments.database, dry_run=not arguments.apply,
    )
    output = report.to_json()
    print(output)
    if arguments.report:
        arguments.report.parent.mkdir(parents=True, exist_ok=True)
        arguments.report.write_text(output + "\n", encoding="utf-8")


if __name__ == "__main__":
    main()
